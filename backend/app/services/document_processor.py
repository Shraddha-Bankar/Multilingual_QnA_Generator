import os
import io
import docx
from pypdf import PdfReader
from typing import Tuple
from app.models.schemas import DocumentMetadata
from app.services.text_cleaner import TextCleaner
from app.services.language_detector import LanguageDetector
from app.config import settings

class DocumentProcessingError(Exception):
    pass

class DocumentProcessor:
    @staticmethod
    def format_file_size(size_in_bytes: int) -> str:
        if size_in_bytes < 1024:
            return f"{size_in_bytes} B"
        elif size_in_bytes < 1024 * 1024:
            return f"{size_in_bytes / 1024:.1f} KB"
        else:
            return f"{size_in_bytes / (1024 * 1024):.2f} MB"

    @classmethod
    def process_file(cls, filename: str, content_bytes: bytes) -> Tuple[str, DocumentMetadata]:
        """
        Extracts text from PDF, DOCX, or TXT file and builds DocumentMetadata.
        Raises DocumentProcessingError if file is unsupported, empty, or unreadable.
        """
        if not content_bytes or len(content_bytes) == 0:
            raise DocumentProcessingError("The uploaded file is empty.")

        file_size_bytes = len(content_bytes)
        max_size_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
        if file_size_bytes > max_size_bytes:
            raise DocumentProcessingError(
                f"File size ({cls.format_file_size(file_size_bytes)}) exceeds the maximum allowed limit of {settings.MAX_UPLOAD_SIZE_MB}MB."
            )

        _, ext = os.path.splitext(filename.lower())
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise DocumentProcessingError(
                f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file."
            )

        raw_text = ""
        page_count = 1

        if ext == ".pdf":
            raw_text, page_count = cls._extract_pdf(content_bytes)
        elif ext == ".docx":
            raw_text, page_count = cls._extract_docx(content_bytes)
        elif ext == ".txt":
            raw_text, page_count = cls._extract_txt(content_bytes)
        else:
            raise DocumentProcessingError(f"Unsupported format: {ext}")

        cleaned_text = TextCleaner.clean_text(raw_text)
        if not cleaned_text or len(cleaned_text.strip()) < 20:
            if ext == ".pdf":
                raise DocumentProcessingError(
                    "No extractable text was found in this PDF. Please upload a text-based PDF or DOCX/TXT document."
                )
            else:
                raise DocumentProcessingError("No extractable text was found in this document.")

        words = cleaned_text.split()
        word_count = len(words)
        char_count = len(cleaned_text)

        lang_name, lang_code = LanguageDetector.detect_language(cleaned_text)

        metadata = DocumentMetadata(
            filename=filename,
            file_type=ext.replace(".", "").upper(),
            file_size_bytes=file_size_bytes,
            file_size_formatted=cls.format_file_size(file_size_bytes),
            page_count=page_count,
            word_count=word_count,
            character_count=char_count,
            detected_language=lang_name,
            detected_language_code=lang_code,
            chunks_count=max(1, (word_count // settings.CHUNK_SIZE_WORDS) + (1 if word_count % settings.CHUNK_SIZE_WORDS else 0))
        )

        return cleaned_text, metadata

    @staticmethod
    def _extract_pdf(content_bytes: bytes) -> Tuple[str, int]:
        try:
            stream = io.BytesIO(content_bytes)
            reader = PdfReader(stream)
            
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    raise DocumentProcessingError("The PDF is password protected and could not be read.")

            page_count = len(reader.pages)
            if page_count == 0:
                raise DocumentProcessingError("The PDF file has 0 pages.")

            extracted_parts = []
            for idx, page in enumerate(reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_parts.append(page_text)
                except Exception as pe:
                    # Skip problematic page but continue
                    continue

            full_text = "\n\n".join(extracted_parts)
            return full_text, page_count
        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f"Failed to parse PDF document: {str(e)}")

    @staticmethod
    def _extract_docx(content_bytes: bytes) -> Tuple[str, int]:
        try:
            stream = io.BytesIO(content_bytes)
            doc = docx.Document(stream)
            
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract text from tables if any
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_texts.append(" | ".join(row_cells))
            
            all_parts = paragraphs + table_texts
            full_text = "\n\n".join(all_parts)
            # Estimate pages (~300 words per page for DOCX)
            word_count = len(full_text.split())
            estimated_pages = max(1, (word_count // 300) + 1)
            return full_text, estimated_pages
        except Exception as e:
            raise DocumentProcessingError(f"Failed to parse DOCX document: {str(e)}")

    @staticmethod
    def _extract_txt(content_bytes: bytes) -> Tuple[str, int]:
        encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
        decoded_text = None
        for enc in encodings_to_try:
            try:
                decoded_text = content_bytes.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if decoded_text is None:
            raise DocumentProcessingError("Unable to decode the text file with standard character encodings.")

        words = decoded_text.split()
        estimated_pages = max(1, (len(words) // 400) + 1)
        return decoded_text, estimated_pages
