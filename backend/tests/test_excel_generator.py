import os
import openpyxl
import pytest
from app.models.schemas import QnAPair
from app.services.excel_generator import ExcelGenerator

def test_excel_generator_structure_and_unicode(tmp_path):
    en_pairs = [
        QnAPair(id=1, question="What is Indian Knowledge Systems?", answer="IKS represents the cumulative wisdom of ancient India."),
        QnAPair(id=2, question="Which texts form the core of IKS?", answer="The Vedas, Upanishads, and Vedangas form the core.")
    ]
    hi_pairs = [
        QnAPair(id=1, question="भारतीय ज्ञान प्रणाली (IKS) क्या है?", answer="IKS प्राचीन भारत के संचयी ज्ञान का प्रतिनिधित्व करता है।"),
        QnAPair(id=2, question="IKS का मुख्य आधार क्या है?", answer="वेद, उपनिषद और वेदांग इसका मुख्य आधार हैं।")
    ]
    mr_pairs = [
        QnAPair(id=1, question="भारतीय ज्ञान प्रणाली (IKS) काय आहे?", answer="IKS प्राचीन भारताच्या संचित ज्ञानाचे प्रतिनिधित्व करते."),
        QnAPair(id=2, question="IKS चा मुख्य आधार काय आहे?", answer="वेद, उपनिषदे आणि वेदांग हा मुख्य पाया आहे.")
    ]

    excel_file = os.path.join(tmp_path, "Test_Multilingual_QnA.xlsx")
    ExcelGenerator.generate_multilingual_workbook(en_pairs, hi_pairs, mr_pairs, excel_file)

    assert os.path.exists(excel_file)

    wb = openpyxl.load_workbook(excel_file)
    
    # 1. Exactly 3 worksheets
    sheet_names = wb.sheetnames
    assert len(sheet_names) == 3
    assert sheet_names == ["English", "Hindi", "Marathi"]

    # 2. Check each sheet structure
    for name in ["English", "Hindi", "Marathi"]:
        ws = wb[name]
        # Column headers
        assert ws["A1"].value == "Questions"
        assert ws["B1"].value == "Answers"
        # Only 2 columns
        assert ws.max_column == 2
        # 2 data rows + 1 header row = 3 rows
        assert ws.max_row == 3

    # Check Unicode content in Hindi and Marathi sheets
    assert "भारतीय ज्ञान प्रणाली" in wb["Hindi"]["A2"].value
    assert "काय आहे" in wb["Marathi"]["A2"].value
