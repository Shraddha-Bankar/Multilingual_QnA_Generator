import os
import io
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

IKS_CONTENT = """Indian Knowledge Systems (IKS): Foundations, Philosophy, and Contributions

1. Overview and Core Philosophy
Indian Knowledge Systems (IKS) encompass the vast corpus of traditional knowledge, philosophy, scientific methodologies, mathematics, metallurgy, architecture, astronomy, and health sciences developed on the Indian subcontinent over millennia. At the core of IKS is an integrated worldview that perceives knowledge as holistic, connecting humanity, nature, and the universe.

2. Primary Textual Corpus and Vedic Foundations
The foundation of IKS resides in the Vedic literature: Rigveda, Samaveda, Yajurveda, and Atharvaveda, accompanied by the Aranyakas, Brahmanas, and Upanishads. Complementing these are the six Vedangas: Shiksha (phonetics), Kalpa (rituals and social law), Vyakarana (grammar), Nirukta (etymology), Chandas (prosody), and Jyotisha (astronomy). Panini's Ashtadhyayi stands as a monumental work in formal linguistics and generative grammar.

3. Mathematics and Astronomy
Ancient and medieval Indian mathematicians made pioneering contributions to global science. Aryabhata introduced foundational concepts of algebra, spherical trigonometry, and proposed that the Earth rotates on its axis. Brahmagupta established the mathematical principles of zero (Shunya) as a number and rules for negative integers. Bhaskaracharya elaborated on calculus principles, planetary positions, and the cyclic method (Chakravala) for solving indeterminate quadratic equations.

4. Health Sciences: Ayurveda and Yoga
Ayurveda, detailed in classic treatises like Charaka Samhita and Sushruta Samhita, emphasizes preventive health, balancing the three doshas (Vata, Pitta, Kapha), and holistic well-being. Sushruta is celebrated for pioneering surgical techniques, including rhinoplasty and cataract surgeries. Patanjali's Yoga Sutras provide a systematic eight-limbed framework (Ashtanga Yoga) for mental discipline, focus, and spiritual development.

5. Metallurgy, Architecture, and Sustainable Practices
Historical Indian metallurgy achieved extraordinary feats such as the rust-resistant Iron Pillar of Delhi and the casting of high-grade Wootz steel. Vastu Shastra and temple architecture reflected sophisticated geometric planning, acoustic engineering, and environmental harmony. Agricultural and water management practices, such as stepwells (Baolis) and Ahar-Pyne systems, demonstrated sustainable ecological stewardship.
"""

ML_CONTENT = """Machine Learning and Modern Artificial Intelligence: Core Principles

1. Introduction to Machine Learning
Machine Learning (ML) is a branch of artificial intelligence focused on building systems that learn from data, identify patterns, and make decisions with minimal human intervention. Rather than following explicitly programmed rules, ML algorithms optimize mathematical objective functions to generalize from past observations.

2. Paradigms of Learning
Machine learning methodologies are categorized into three primary paradigms:
- Supervised Learning: Algorithms learn a mapping function from input features to known target labels using labeled training datasets. Common algorithms include Linear Regression, Support Vector Machines, Random Forests, and Gradient Boosting.
- Unsupervised Learning: Systems discover hidden structures, groupings, or representations within unlabeled data. Prominent techniques include K-Means Clustering, Principal Component Analysis (PCA), and Autoencoders.
- Reinforcement Learning: Agents interact with dynamic environments to maximize cumulative numerical rewards through trial-and-error exploration and exploitation, widely applied in robotics, gaming, and autonomous navigation.

3. Deep Learning and Neural Network Architectures
Deep Learning utilizes deep artificial neural networks inspired by biological neural circuits. Convolutional Neural Networks (CNNs) specialize in grid-like visual data for computer vision. Recurrent Neural Networks (RNNs) and Transformers revolutionize natural language processing through multi-head self-attention mechanisms that capture long-range contextual relationships.

4. Model Evaluation and Generalization
Evaluating machine learning models requires rigorous validation methodologies to prevent overfitting. Metrics such as Precision, Recall, F1-Score, ROC-AUC, and Mean Squared Error assess model generalization on unseen test distributions. Regularization techniques like Dropout, L1/L2 penalties, and cross-validation safeguard model stability.
"""

def generate_samples():
    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_inputs")
    os.makedirs(out_dir, exist_ok=True)

    # 1. TXT files
    with open(os.path.join(out_dir, "Indian_Knowledge_Systems.txt"), "w", encoding="utf-8") as f:
        f.write(IKS_CONTENT)
    with open(os.path.join(out_dir, "Machine_Learning_Fundamentals.txt"), "w", encoding="utf-8") as f:
        f.write(ML_CONTENT)

    # 2. DOCX files
    doc_iks = docx.Document()
    for para in IKS_CONTENT.split("\n\n"):
        if para.strip():
            doc_iks.add_paragraph(para.strip())
    doc_iks.save(os.path.join(out_dir, "Indian_Knowledge_Systems.docx"))

    doc_ml = docx.Document()
    for para in ML_CONTENT.split("\n\n"):
        if para.strip():
            doc_ml.add_paragraph(para.strip())
    doc_ml.save(os.path.join(out_dir, "Machine_Learning_Fundamentals.docx"))

    # 3. PDF files
    styles = getSampleStyleSheet()
    style_h = styles['Heading1']
    style_b = styles['Normal']
    style_b.fontSize = 10
    style_b.leading = 14

    def build_pdf(text, path):
        doc = SimpleDocTemplate(path, pagesize=letter)
        story = []
        lines = text.split("\n\n")
        for i, block in enumerate(lines):
            if i == 0:
                story.append(Paragraph(block, style_h))
            else:
                story.append(Paragraph(block.replace("\n", "<br/>"), style_b))
            story.append(Spacer(1, 10))
        doc.build(story)

    build_pdf(IKS_CONTENT, os.path.join(out_dir, "Indian_Knowledge_Systems.pdf"))
    build_pdf(ML_CONTENT, os.path.join(out_dir, "Machine_Learning_Fundamentals.pdf"))

    print("Sample files successfully generated in sample_inputs/")

if __name__ == "__main__":
    generate_samples()
